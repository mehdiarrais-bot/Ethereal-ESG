"""Nombres à la française dans les livrables (DETTE.md § 5, audit 2026-09-24).

Avant : « 21,500 MWh » (ancien générateur) et « 21 500 MWh » (texte
analytique) dans le même rapport. Le test lit le texte RENDU de chaque
livrable français — pas le contenu intermédiaire — et refuse tout nombre
à l'anglaise ; en anglais, rien ne doit changer.
"""
import io
import os
import re
import sys
from typing import cast

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from test_suite import make_request                    # noqa: E402
from esg_calculator import calculate_esg_scores        # noqa: E402
from content_generator import generate_esg_content     # noqa: E402
from typo import fr_numbers, fix, language, NBSP       # noqa: E402

ANGLAIS = re.compile(r"\d,\d{3}(?!\d)|\d\.\d")


@pytest.mark.parametrize("avant,apres", [
    ("21,500 MWh", f"21{NBSP}500 MWh"),
    ("1,234,567 t", f"1{NBSP}234{NBSP}567 t"),
    ("TF 6.2", "TF 6,2"),
    ("8,200.5 t", f"8{NBSP}200,5 t"),
    ("42% et 50 %", f"42{NBSP}% et 50{NBSP}%"),
    ("Scope 1, 2 et 3", "Scope 1, 2 et 3"),          # énumération : intacte
    ("1. Synthèse Exécutive", "1. Synthèse Exécutive"),  # numérotation : intacte
    ("Règlement (UE) 2020/852", "Règlement (UE) 2020/852"),
])
def test_fr_numbers(avant, apres):
    assert fr_numbers(avant) == apres


def test_anglais_intact():
    with language("en"):
        assert fix("21,500 MWh, 6.2, 42%") == "21,500 MWh, 6.2, 42%"
    assert fix("21,500") == "21,500"          # hors composition : identité


def _pdf_text(data):
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    return " ".join(str(doc[i].get_text()) for i in range(doc.page_count))


def _docx_text(data):
    from docx import Document
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return " ".join(parts)


def _pptx_text(data):
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape
    out = []
    for slide in Presentation(io.BytesIO(data)).slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(cast(Shape, sh).text_frame.text)
    return " ".join(out)


def _livrables(lang):
    from report_generator import generate_pdf_report
    from onepager_generator import generate_onepager_pdf
    from docx_generator import generate_word_report
    from proposal_generator import generate_proposal_docx
    from ppt_generator import generate_pptx
    r = make_request(lang, environmental=make_request().environmental.model_copy(
        update={"energy_consumption_mwh": 21500, "water_consumption_m3": 120000}))
    s = calculate_esg_scores(r)
    c = generate_esg_content(r, s)
    return {
        "pdf": _pdf_text(generate_pdf_report(r, s, c, {})),
        "onepager": _pdf_text(generate_onepager_pdf(r, s)),
        "docx": _docx_text(generate_word_report(r, s, c)),
        "proposal": _docx_text(generate_proposal_docx(r, s)),
        "pptx": _pptx_text(generate_pptx(r, s, c, {})),
    }


def test_aucun_nombre_a_l_anglaise_dans_les_livrables_francais():
    for nom, texte in _livrables("fr").items():
        fautes = sorted(set(m.group(0) for m in ANGLAIS.finditer(texte)))
        assert not fautes, f"{nom} : {fautes}"


def test_livrables_anglais_inchanges():
    textes = _livrables("en")
    assert "21,500" in textes["pdf"] or "21,500" in textes["docx"]
    assert NBSP + "%" not in textes["docx"]
