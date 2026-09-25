"""Diagnostic d'ensemble et conclusion propres au dossier (synthesis.py).

Lot 2 du chantier « textes » (2026-09-26) : la synthèse et la conclusion
changent avec les données — enjeux classés par gravité, liens entre piliers,
« si rien ne change », 90 premiers jours.
"""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from test_suite import (make_request, ENGAGEMENTS_FABRIQUES, ATTRIBUTIONS_LEGALES_FAUSSES,  # noqa: E402
                        MARQUEURS_DE_PROVENANCE, AFFIRMATIONS_INTERDITES, MARQUEURS_METHODO_INVENTEE)
from test_affirmations_reglementaires import CSRD_OBLIGATION_INCONDITIONNELLE  # noqa: E402
from test_analyse_approfondie import _trois, _tous_les_textes                   # noqa: E402
from esg_calculator import calculate_esg_scores                                 # noqa: E402
from models import GovernanceData, SocialData, EnvironmentalData                 # noqa: E402
import synthesis as SY                                                           # noqa: E402
import synthesis_texts as STX                                                    # noqa: E402


def _cles(r):
    return [i.key for i in SY.issues(r, calculate_esg_scores(r))]


def test_le_premier_enjeu_change_avec_l_entreprise():
    premiers = {n: _cles(r)[0] for n, r in _trois().items()}
    assert premiers == {"transport": "safety", "services": "retention", "industrie": "integrity"}


def test_enjeux_classes_par_gravite():
    r = _trois()["transport"]
    sev = [i.severity for i in SY.issues(r, calculate_esg_scores(r))]
    assert sev == sorted(sev, reverse=True)


def test_aucun_enjeu_sans_sa_donnee():
    r = make_request(social=SocialData(total_employees=10),
                     governance=GovernanceData(board_members=5),
                     environmental=EnvironmentalData(co2_emissions_tonnes=10))
    cles = set(_cles(r))
    assert not cles & {"safety", "retention", "skills", "mix", "steering", "independence",
                       "integrity", "renewable", "waste"}, cles


def test_booleen_inconnu_n_est_pas_un_non():
    r = make_request(governance=GovernanceData(independent_board_percent=45))
    assert "steering" not in _cles(r)
    r = make_request(governance=GovernanceData(sustainability_committee=True,
                                               independent_board_percent=45))
    assert "steering" not in _cles(r)


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_sorties_completes_et_remplies(lang):
    for r in _trois(lang).values():
        s = calculate_esg_scores(r)
        o, c = SY.overview(r, s), SY.closing(r, s)
        texte = " ".join(o["profile"] + o["links"] + c["retain"] + c["horizon"] + c["first"]
                         + [x for i in o["issues"] for x in (i.title, i.cause, i.csq)])
        assert "{" not in texte and "}" not in texte and "None" not in texte
        assert len(c["first"]) == min(SY.TOP, len(SY.issues(r, s)))


def test_elision_du_nom():
    r = make_request()          # « Acme Industries »
    profil = SY.overview(r, calculate_esg_scores(r))["profile"][0]
    assert "d'Acme Industries" in profil and "de Acme" not in profil


def test_liens_entre_piliers_du_transporteur():
    r = _trois()["transport"]
    liens = " ".join(SY.overview(r, calculate_esg_scores(r))["links"])
    assert "n'a pas d'instance de gouvernance" in liens      # sinistralité × pas de comité
    assert "même chantier de ressources humaines" in liens    # rotation × formation


def test_parite_fr_en():
    for fr, en in ((STX.ISSUES["fr"], STX.ISSUES["en"]), (STX.STRENGTHS["fr"], STX.STRENGTHS["en"]),
                   (STX.LINKS["fr"], STX.LINKS["en"]), (STX.T["fr"], STX.T["en"])):
        assert fr.keys() == en.keys()
    for k in STX.ISSUES["fr"]:
        assert STX.ISSUES["fr"][k].keys() == STX.ISSUES["en"][k].keys(), k


def test_aucune_affirmation_interdite():
    blob = " ".join(str(v) for v in (STX.ISSUES, STX.STRENGTHS, STX.LINKS, STX.T)).lower()
    blob += _tous_les_textes()
    for m in (ENGAGEMENTS_FABRIQUES + ATTRIBUTIONS_LEGALES_FAUSSES + MARQUEURS_DE_PROVENANCE
              + AFFIRMATIONS_INTERDITES + MARQUEURS_METHODO_INVENTEE
              + CSRD_OBLIGATION_INCONDITIONNELLE):
        assert m.lower() not in blob, m
    assert "non une obligation légale" in STX.ISSUES["fr"]["mix"]["csq"]


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_dans_le_pdf_et_le_word(lang):
    from test_parite import _pdf, _docx, _norm
    from content_generator import generate_esg_content
    from report_generator import generate_pdf_report
    from docx_generator import generate_word_report
    r = _trois(lang)["transport"]
    s = calculate_esg_scores(r)
    c = generate_esg_content(r, s)
    attendus = [STX.T[lang]["overview"], STX.T[lang]["first_title"],
                SY.issues(r, s)[0].title]
    pdf = _norm(_pdf(generate_pdf_report(r, s, c, {})), lang)
    docx = _norm(_docx(generate_word_report(r, s, c, charts={})), lang)
    for a in attendus:
        assert _norm(a, lang) in pdf, a
        assert _norm(a, lang) in docx, a


# ── Lot 3 : présentation et synthèse une page ─────────────────────────────

def _slides_text(data) -> list[str]:
    import io
    from pptx import Presentation
    out = []
    for sl in Presentation(io.BytesIO(data)).slides:
        out.append(" ".join(getattr(sh, "text_frame").text for sh in sl.shapes
                            if sh.has_text_frame))
    return out


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_presentation_enjeux_analyses_et_90_jours(lang):
    from content_generator import generate_esg_content
    from ppt_generator import generate_pptx
    import analysis as AN
    import illustrations as IL
    r = _trois(lang)["transport"]
    s = calculate_esg_scores(r)
    slides = _slides_text(generate_pptx(r, s, generate_esg_content(r, s), IL.build(r)))
    joined = " ".join(slides)
    assert STX.T[lang]["issues_title"] in joined
    assert STX.T[lang]["first_title"] in joined
    assert SY.issues(r, s)[0].title in joined
    for pillar in ("env", "social", "gov"):
        assert AN.analysis_title(r, pillar) in joined, pillar


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_synthese_une_page_porte_le_diagnostic(lang):
    from test_parite import _pdf, _norm
    from onepager_generator import generate_onepager_pdf
    r = _trois(lang)["transport"]
    s = calculate_esg_scores(r)
    texte = _norm(_pdf(generate_onepager_pdf(r, s)), lang)
    assert _norm(STX.T[lang]["overview"].upper(), lang) in texte
    assert _norm(SY.issues(r, s)[0].title, lang) in texte
