"""Même contenu analytique dans le PDF, le Word et la présentation.

Audit du 2026-09-24 : le README promettait « Rapport Word — même contenu,
éditable », mais le texte analytique (narrative.py) n'existait que dans le
PDF. Chaque phrase clé doit se retrouver dans les trois livrables — dans les
notes de l'orateur pour la présentation.
"""
import io
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from test_suite import make_request                    # noqa: E402
from esg_calculator import calculate_esg_scores        # noqa: E402
from content_generator import (generate_esg_content, risks_opportunities,  # noqa: E402
                               compliance_assessment, enriched_recommendations, roadmap_12m)
import narrative as NR                                  # noqa: E402
from typo import fr_numbers                             # noqa: E402


def _norm(text: str, lang: str) -> str:
    text = fr_numbers(text) if lang == "fr" else text
    return re.sub(r"\s+", " ", text.replace(" ", " ")).strip()


def _pdf(data):
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    return " ".join(str(doc[i].get_text()) for i in range(doc.page_count))


def _docx(data):
    from docx import Document
    return " ".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def _pptx_notes(data):
    from pptx import Presentation
    return " ".join(s.notes_slide.notes_text_frame.text
                    for s in Presentation(io.BytesIO(data)).slides
                    if s.has_notes_slide and s.notes_slide.notes_text_frame is not None)


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_texte_analytique_dans_les_trois_livrables(lang):
    from report_generator import generate_pdf_report
    from docx_generator import generate_word_report
    from ppt_generator import generate_pptx
    r = make_request(lang)
    s = calculate_esg_scores(r)
    c = generate_esg_content(r, s)
    ro, gaps = risks_opportunities(r, s), compliance_assessment(r, s)
    recs, rm = enriched_recommendations(r, s), roadmap_12m(r, s)
    ref = "CSRD / ESRS"
    phrases = {
        "présentation": NR.company_paragraphs(r, ref)[0],
        "lecture environnement": NR.pillar_paragraphs(r, s, "env")[0],
        "lecture gouvernance": NR.pillar_paragraphs(r, s, "gov")[1],
        "leviers social": NR.levers_sentence(r, recs, "social"),
        "empreinte carbone": NR.ghg_paragraph(r),
        "chapitre diagnostic": NR.act1_intro(r, s, gaps, ro["risks"]),
        "risques": NR.risks_intro(r, ro["risks"]),
        "plan d'action": NR.recs_intro(r, recs),
        "feuille de route": NR.roadmap_intro(r, rm),
    }
    livrables = {
        "pdf": _pdf(generate_pdf_report(r, s, c, {})),
        "docx": _docx(generate_word_report(r, s, c)),
        "pptx (notes)": _pptx_notes(generate_pptx(r, s, c, {})),
    }
    manques = [(nom, cle) for nom, texte in livrables.items()
               for cle, phrase in phrases.items()
               if _norm(phrase, lang) not in _norm(texte, lang)]
    assert not manques, manques
